import os
import json
import httpx
import yaml
from openpyxl import load_workbook
from dotenv import load_dotenv

load_dotenv()

class AIService:
    def __init__(self):
        # We'll fetch it here, but also have a fallback in the method
        self.api_key = os.getenv("GROQ_API_KEY")

    def _read_extracted_files(self, attachments):
        """Reads JSON and Excel files from the extracted zip attachments."""
        file_contents = ""
        for att in attachments:
            extracted_dir = att.get("extracted_to")
            if not extracted_dir:
                continue
                
            for root, _, files in os.walk(extracted_dir):
                for file in files:
                    file_path = os.path.join(root, file)
                    # Skip hidden macOS files
                    if file.startswith("._") or file == ".DS_Store":
                        continue
                        
                    if file.endswith('.json'):
                        try:
                            with open(file_path, 'r', encoding='utf-8') as f:
                                content = f.read()
                                file_contents += f"\n\n--- Content of {file} ---\n{content}"
                        except Exception as e:
                            print(f"Error reading {file}: {e}")
                    elif file.endswith('.xlsx'):
                        try:
                            wb = load_workbook(filename=file_path, data_only=True)
                            file_contents += f"\n\n--- Content of {file} ---\n"
                            for sheetname in wb.sheetnames:
                                sheet = wb[sheetname]
                                file_contents += f"Sheet: {sheetname}\n"
                                for row in sheet.iter_rows(values_only=True):
                                    row_str = " | ".join([str(cell) if cell is not None else "" for cell in row])
                                    file_contents += f"{row_str}\n"
                        except Exception as e:
                            print(f"Error reading {file}: {e}")
        return file_contents

    async def generate_oas(self, context: dict):
        """
        Calls Groq API to generate an OAS using the collected Jira context.
        """
        issue_key = context.get("issue_key")
        print(f"🤖 AI Service: Starting OAS generation for {issue_key} using Groq API...")
        
        # Ensure API key is set (fallback fetch if init failed)
        if not self.api_key:
            self.api_key = os.getenv("GROQ_API_KEY")

        if not self.api_key:
            raise ValueError("GROQ_API_KEY is not set in the environment variables.")

        # Gather extracted file content
        extracted_files_content = self._read_extracted_files(context.get("attachments", []))
        
        # Build the prompt
        system_prompt = (
            "You are an expert OpenAPI Specification (OAS) generator. "
            "Generate a valid YAML OpenAPI 3.1 specification based on the provided Jira issue context and attached schemas/excel data. "
            "IMPORTANT: All schema definitions MUST be included inside the 'components/schemas' section of the YAML. "
            "DO NOT use external $ref pointers to .json or .yaml files. "
            "All $ref pointers MUST be internal, e.g., '$ref: \"#/components/schemas/MySchema\"'. "
            "Return ONLY the raw YAML string inside your response, without any markdown formatting code blocks. Do NOT include ```yaml tags."
        )
        
        user_prompt = f"""
        **Jira Issue Key:** {issue_key}
        **Summary:** {context.get('summary')}
        **Description:** {context.get('description')}
        **Comments:** {json.dumps(context.get('comments', []), indent=2)}
        
        **Attached Files (Excel Logic / JSON Samples):**
        {extracted_files_content}
        
        Please generate the OpenAPI YAML specification representing this API definition. 
        It should accurately consolidate the rules from the description and the attached schemas.
        """

        headers = {
            "Authorization": f"Bearer {self.api_key}",
            "Content-Type": "application/json"
        }
        
        payload = {
            "model": "llama-3.3-70b-versatile",
            "messages": [
                {"role": "system", "content": system_prompt},
                {"role": "user", "content": user_prompt}
            ],
            "temperature": 0.2
        }

        try:
            async with httpx.AsyncClient(timeout=120.0) as client:
                response = await client.post(
                    "https://api.groq.com/openai/v1/chat/completions",
                    headers=headers,
                    json=payload
                )
                response.raise_for_status()
                response_data = response.json()
                
                ai_text = response_data['choices'][0]['message']['content'].strip()
                
                # Basic cleanup of markdown blocks
                for tag in ["```yaml", "```yml", "```json", "```"]:
                    if ai_text.startswith(tag):
                        ai_text = ai_text[len(tag):]
                if ai_text.endswith("```"):
                    ai_text = ai_text[:-3]
                    
                ai_text = ai_text.strip()
                
                # Verify it's valid YAML
                oas_spec = yaml.safe_load(ai_text)
                
                print(f"✅ AI Service: Generation complete (YAML) for {issue_key}")
                
                return {
                    "status": "success",
                    "issue_key": issue_key,
                    "oas_spec": oas_spec,
                    "raw_text": ai_text
                }
                
        except yaml.YAMLError as ye:
            failed_text = locals().get('ai_text', '')[:500]
            print(f"❌ AI Service YAML Error: {ye}")
            raise Exception(f"AI response was not a valid YAML structure. Response text was: {failed_text}")
        except Exception as e:
            print(f"❌ AI Service Error: {e}")
            raise e

ai_service = AIService()
